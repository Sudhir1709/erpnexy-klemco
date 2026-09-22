"""parse_items_file — the server half of the Excel/CSV "Import Items" button (Sales Order + Quotation).

Column mapping is client-side (public/js/item_import.js); the server only has to hand back the rows
of a .csv / .xlsx File the caller may read.
"""
import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils.xlsxutils import make_xlsx

from klemco_cs.item_import import parse_items_file


class TestItemImport(FrappeTestCase):
    def _file(self, file_name, content):
        doc = frappe.get_doc({
            "doctype": "File", "file_name": file_name, "is_private": 1, "content": content,
        }).insert()
        # The DB rolls back after each test, the file on disk does not — delete explicitly.
        self.addCleanup(lambda: frappe.delete_doc("File", doc.name, force=True, ignore_permissions=True))
        return doc

    def test_csv_rows(self):
        f = self._file("klemco_test_items.csv", b"\xef\xbb\xbfItem Code,Qty,Rate\nKL-001,2,10\n")
        rows = parse_items_file(f.file_url)
        self.assertEqual(rows[0], ["Item Code", "Qty", "Rate"])   # BOM stripped
        self.assertEqual(rows[1], ["KL-001", "2", "10"])

    def test_xlsx_rows(self):
        content = make_xlsx([["Item Code", "Qty"], ["KL-001", 2]], "Items").getvalue()
        f = self._file("klemco_test_items.xlsx", content)
        rows = parse_items_file(f.file_url)
        self.assertEqual(rows[0], ["Item Code", "Qty"])
        self.assertEqual(rows[1][0], "KL-001")
        self.assertEqual(rows[1][1], 2)   # native cell types

    def test_rejects_other_extensions(self):
        self.assertRaises(frappe.ValidationError, parse_items_file, "/private/files/items.xls")
        self.assertRaises(frappe.ValidationError, parse_items_file, "/private/files/items.pdf")

    def test_download_template(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from klemco_cs.item_import import download_template

        for doctype, expected in (
            ("Quotation", ["Item Code", "Qty", "Rate", "Warehouse"]),
            ("Sales Order", ["Item Code", "Qty", "Rate", "Warehouse", "Delivery Date"]),
        ):
            frappe.response.pop("filecontent", None)
            download_template(doctype)
            self.assertTrue(frappe.response["filename"].endswith(".xlsx"))
            self.assertEqual(frappe.response["type"], "binary")
            wb = load_workbook(BytesIO(frappe.response["filecontent"]))
            self.assertEqual(wb.sheetnames, ["Items", "How to"])
            header = [c.value for c in next(wb["Items"].iter_rows())]
            self.assertEqual(header, expected, doctype)
            self.assertEqual(wb["Items"].max_row, 1, "template must have no data rows")
        self.assertRaises(frappe.ValidationError, download_template, "Item")

    def test_requires_read_permission_on_the_file(self):
        f = self._file("klemco_test_private.csv", b"Item Code,Qty\nKL-001,1\n")
        frappe.set_user("Guest")
        self.addCleanup(frappe.set_user, "Administrator")
        self.assertRaises(frappe.PermissionError, parse_items_file, f.file_url)
