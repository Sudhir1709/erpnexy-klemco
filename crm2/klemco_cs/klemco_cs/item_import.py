# Parse an uploaded Excel (.xlsx) or CSV into rows for the "Import Items" button on Sales Order and
# Quotation (public/js/item_import.js), so a user can upload a plain spreadsheet of Item Code + Qty
# instead of the grid's 7-row template. Column mapping happens client-side; this only returns rows.

import frappe
from frappe import _

MAX_ROWS = 5000   # same ceiling as Frappe's stock grid upload


@frappe.whitelist()
def parse_items_file(file_url):
    """Return the file's rows as a list of lists (header row included). Supports .xlsx and .csv.
    The File record must be readable by the caller — normally their own fresh upload."""
    if not file_url:
        return []
    lname = file_url.lower()
    if not lname.endswith((".csv", ".xlsx")):
        frappe.throw(_("Upload a .csv or .xlsx file (legacy .xls is not supported — save it as .xlsx)."))

    fdoc = _get_readable_file(file_url)
    if lname.endswith(".xlsx"):
        from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file
        rows = read_xlsx_file_from_attached_file(filepath=fdoc.get_full_path())
    else:
        from frappe.utils.csvutils import read_csv_content
        content = fdoc.get_content()
        if isinstance(content, bytes):
            content = content.decode("utf-8-sig", errors="replace")
        rows = read_csv_content(content.lstrip("﻿"))

    rows = rows or []
    if len(rows) > MAX_ROWS:
        frappe.throw(_("Cannot import a file with more than {0} rows.").format(MAX_ROWS))
    return rows


TEMPLATE_DOCTYPES = ("Quotation", "Sales Order")


@frappe.whitelist()
def download_template(doctype):
    """Excel template for the Import Items upload: sheet 'Items' holds the header row the importer
    expects (Delivery Date only where the child table has it), sheet 'How to' the short instructions.
    No example data row — it would import as a bogus item code."""
    if doctype not in TEMPLATE_DOCTYPES:
        frappe.throw(_("No item import template for {0}.").format(doctype))
    frappe.has_permission(doctype, "read", throw=True)
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    child_dt = frappe.get_meta(doctype).get_field("items").options
    columns = ["Item Code", "Qty", "Rate", "Warehouse"]
    if frappe.get_meta(child_dt).has_field("delivery_date"):
        columns.append("Delivery Date")

    how_to = [
        ["How to fill the 'Items' sheet"],
        ["Item Code — required; the ERPNext Item Code exactly as on the Item master."],
        ["Qty — quantity (blank = 1)."],
        ["Rate — optional; leave blank to take the price list rate."],
        ["Warehouse — optional; leave blank for the default."],
    ]
    if "Delivery Date" in columns:
        how_to.append(["Delivery Date — optional, per line; dd.mm.yyyy or yyyy-mm-dd."])
    how_to += [
        ["Keep the header row; extra columns are ignored."],
        [f"Upload the file on the {doctype} with the items grid's 'Upload' button (or Get Items From → Import Items)."],
        ["Rows are added to the grid; Item Name, UOM and prices fill in automatically."],
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for idx in range(1, len(columns) + 1):   # not `_` — that would shadow the translator
        ws.column_dimensions[get_column_letter(idx)].width = 18
    ws2 = wb.create_sheet("How to")
    for row in how_to:
        ws2.append(row)
    ws2["A1"].font = Font(bold=True)
    ws2.column_dimensions["A"].width = 110
    buf = BytesIO()
    wb.save(buf)
    frappe.response["filename"] = f"{doctype} Items Template.xlsx"
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"


def _get_readable_file(file_url):
    """One file_url can belong to several File records (content de-duplication) — return one the
    current user may read, else refuse. Stops a logged-in user reading arbitrary private files by URL."""
    for name in frappe.get_all("File", filters={"file_url": file_url}, pluck="name"):
        fdoc = frappe.get_doc("File", name)
        if fdoc.has_permission("read"):
            return fdoc
    frappe.throw(_("You do not have access to this file."), frappe.PermissionError)
